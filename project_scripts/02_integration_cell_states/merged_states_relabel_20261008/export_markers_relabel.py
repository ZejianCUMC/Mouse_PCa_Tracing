"""Export the descriptive marker workbook with the approved 2026-10-08 names and re-index (copy of scripts/export_state_markers_excel.py).
Only labels/order change: c6->Partial-EMT(Lum), old c9->c7 Partial-EMT(Mes), old c7->c9 Mes-Fibroblast-like (STATE_MAP.csv).
Same per-state full-gene profiles (profiles/cluster_<raw>_allgenes.csv), same marker rule. Conda env: moscot."""
from pathlib import Path
import json,hashlib,math
import pandas as pd,numpy as np
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule,CellIsRule
from openpyxl.utils import get_column_letter
N=Path(__file__).resolve().parent;d=N.parent;p=d.parents[5]
assert (p/'CURRENT_COHORT.json').exists()
target=p/'Npp_tumor_merged_10states_enriched_markers_PartialEMT.xlsx'
assert not target.exists(),'Preserve existing workbook; choose a new version explicitly'
states=pd.read_csv(d/'state_summary.csv');smap=pd.read_csv(N/'STATE_MAP.csv').set_index('old_cluster')
assert set(states.cluster)==set(smap.index)
states['putative_name']=states.cluster.map(smap.new_name);states['order']=states.cluster.map(smap.order);states['cluster']=states.cluster.map(smap.new_cluster)
states=states.sort_values('order').reset_index(drop=True);profile_summary=pd.read_csv(d/'profiles/cluster_summary.csv').set_index('cluster')
meta=pd.read_csv(d/'final_metadata_PRIVATE.csv',index_col=0)
assert len(meta)==55133 and states.cells.sum()==55133 and len(states)==10
assert states.set_index('cluster').putative_name[['c6','c7','c9']].tolist()==['Partial-EMT(Lum)','Partial-EMT(Mes)','Mes-Fibroblast-like']
rows=[];inputs={};sourcecols=None
for st in states.itertuples():
 source=d/f'profiles/cluster_{st.raw_cluster}_allgenes.csv';g=pd.read_csv(source)
 inputs[source.name]=hashlib.file_digest(source.open('rb'),'sha256').hexdigest()
 assert g.gene.is_unique and len(g)==33648
 g=g[g.gene_id.astype(str).str.startswith('ENSMUSG') & (g.pct_in>=10)&(g.log2fc_mean>=.5)&(g.pct_in>g.pct_out)].copy()
 g['cluster_id']=st.cluster;g['cell_state']=st.cluster+'_'+st.putative_name;g['state_cells']=st.cells
 g['detection_delta_pct_points']=g.pct_in-g.pct_out
 g['sample_consistency']=np.where(g.eligible_samples<6,'Limited samples',np.where((g.sample_positive_fraction>=.7)&(g.sample_median_log2fc>=.5),'Consistent','Mixed/weak'))
 g['gene_category']=np.select([g.gene.str.match(r'^mt-'),g.gene.str.match(r'^Rp[sl]'),g.gene.str.match(r'^Gm\d')],['Mitochondrial','Ribosomal','Predicted Gm'],default='Other endogenous')
 g['dominant_sample']=profile_summary.loc[st.raw_cluster,'dominant_sample'];g['dominant_sample_pct']=profile_summary.loc[st.raw_cluster,'dominant_sample_pct']
 g['cohort_caution']=np.where(g.dominant_sample_pct>=80,'Sample-dominated','')
 g=g.sort_values(['marker_rank_score','log2fc_mean','gene'],ascending=[False,False,True]);g['rank']=np.arange(1,len(g)+1);rows.append(g)
allgenes=pd.concat(rows,ignore_index=True)
enriched=allgenes.groupby('gene').cluster_id.agg(lambda v:','.join(v))
n_enriched=allgenes.groupby('gene').cluster_id.nunique()
allgenes['n_states_enriched']=allgenes.gene.map(n_enriched)
allgenes['enriched_in_states']=allgenes.gene.map(enriched)
allgenes['specificity_flag']=np.where(allgenes.n_states_enriched.eq(1),'One-state enriched','Shared enrichment')
best=allgenes.loc[allgenes.groupby('gene').marker_rank_score.idxmax()].set_index('gene').cluster_id
allgenes['highest_marker_score_state']=allgenes.gene.map(best)
columns=['cluster_id','cell_state','rank','gene','gene_id','state_cells','pct_in','pct_out','detection_delta_pct_points','log2fc_mean','mean_norm','mean_rest','marker_rank_score','n_states_enriched','enriched_in_states','specificity_flag','highest_marker_score_state','eligible_samples','sample_median_log2fc','sample_positive_fraction','sample_consistency','Npp53_eligible_samples','Npp53_median_log2fc','Npp53Rb1_eligible_samples','Npp53Rb1_median_log2fc','gene_category','dominant_sample','dominant_sample_pct','cohort_caution']
allgenes=allgenes[columns];top=allgenes[allgenes['rank']<=30].copy()
summary=[]
for st in states.itertuples():
 g=allgenes[allgenes.cluster_id.eq(st.cluster)];prof=profile_summary.loc[st.raw_cluster]
 summary.append(dict(cluster_id=st.cluster,cell_state=st.cluster+'_'+st.putative_name,cells=st.cells,Npp53=st.Npp53,Npp53Rb1=st.Npp53Rb1,enriched_genes=len(g),one_state_enriched=int(g.n_states_enriched.eq(1).sum()),consistent_markers=int(g.sample_consistency.eq('Consistent').sum()),eligible_samples=st.eligible_samples,dominant_sample=prof.dominant_sample,dominant_sample_pct=prof.dominant_sample_pct,caution='Limited samples; descriptive only' if st.eligible_samples<6 else ('Sample-dominated' if prof.dominant_sample_pct>=80 else ''),top10_markers=', '.join(g.gene.head(10))))
summary=pd.DataFrame(summary)
guide=[
 ('Scope','Current 55,133-cell, 10-state Npp53 + Npp53Rb1 tumor cohort. Names and order approved 2026-10-08: c6_Partial-EMT(Lum), c7_Partial-EMT(Mes) (previously c9_Mes-OxidativeMetabolism), c9_Mes-Fibroblast-like (previously c7_Mes-CAF-like); other names unchanged.'),
 ('Contents','State_summary; Top30_per_state (ranked candidates); All_enriched; one sheet per state. Every data sheet can be filtered and sorted in Excel.'),
 ('Enriched marker rule','Endogenous ENSMUSG feature; detection in state >=10%; pooled log2 fold-change >=0.5; detection in state greater than in other states. All qualifying genes are retained, not just the plotted markers.'),
 ('Comparator','Each current state versus ALL other current tumor states. Not every pairwise state comparison; not restricted to the same epithelial/mesenchymal family.'),
 ('Gene universe','Existing shared-gene profiles: 33,648 shared features. Transgenes/reporters excluded from this marker export. Source-specific features absent from the shared matrix were not tested.'),
 ('Statistical status','DESCRIPTIVE enrichment, not a new differential-expression test. Gene-level p-values/FDR were not calculated; none are implied. AUCell pathway FDR does not apply to these genes.'),
 ('mean_norm / mean_rest','Mean per-cell raw counts normalized to 10,000 using FULL original library total, inside the state / all other states. These are means before log transformation.'),
 ('log2fc_mean','log2((mean_norm + 0.1)/(mean_rest + 0.1)); a descriptive pooled fold-change, not a model coefficient.'),
 ('pct_in / pct_out','Percent of cells with raw count >0, on a 0–100 scale. detection_delta_pct_points = pct_in - pct_out.'),
 ('marker_rank_score','max(log2fc_mean,0) * max((pct_in-pct_out)/100,0). Ranking favors fold-change AND differential detection; it is not a p-value.'),
 ('n_states_enriched / specificity_flag','Number of current states where the gene meets the same marker rule. One-state enriched is threshold-relative, not proof of exclusive expression.'),
 ('highest_marker_score_state','State with the highest saved marker ranking score among qualifying states; ties resolve by current state order. Different from maximum expression.'),
 ('eligible_samples','Within-sample support requires >=20 cells inside AND outside the state. Sample IDs are the comparison units; independence is assumed because animal-level IDs are unavailable.'),
 ('sample_median_log2fc','Median across eligible samples of log2((state pooled raw-count CPM+1)/(other-state pooled raw-count CPM+1)), using full-library totals. Different normalization/pseudocount from log2fc_mean.'),
 ('sample_positive_fraction','Fraction (0–1) of eligible samples with positive within-sample pseudobulk log2 fold-change.'),
 ('sample_consistency','Consistent: >=6 eligible samples, positive fraction >=0.70, median within-sample log2FC >=0.5. Limited samples: <6. Otherwise Mixed/weak. This flag is NOT statistical significance.'),
 ('Model-specific columns','Number of eligible samples and median pseudobulk log2 fold-change separately for Npp53 / Npp53Rb1. Blank means no eligible samples.'),
 ('Cautions','c7 (Partial-EMT(Mes)) has only 2 eligible samples; 97.5% of its cells come from JZ135. Sample-dominated flags use >=80% of state cells from one sample. These flags do not automatically remove genes.'),
 ('Gene categories','Mitochondrial, ribosomal and predicted Gm genes are retained but flagged for filtering; do not interpret them automatically as lineage markers.'),
 ('No biological change','Existing full-gene profiles were exported; only state names and cluster indices changed (2026-10-08 relabel). No cells, clustering, latent space or UMAP changed.'),
]
wb=Workbook();wb.remove(wb.active);dfs={'Read_me':pd.DataFrame(guide,columns=['Item','Explanation']),'State_summary':summary,'Top30_per_state':top,'All_enriched':allgenes}
for st in states.itertuples():dfs[(st.cluster+'_'+st.putative_name)[:31]]=allgenes[allgenes.cluster_id.eq(st.cluster)].copy()
for i,(name,df) in enumerate(dfs.items()):
 ws=wb.create_sheet(name);ws.append(list(df.columns))
 for row in df.itertuples(index=False,name=None):
  ws.append([None if pd.isna(v) else (v.item() if isinstance(v,np.generic) else v) for v in row])
 for row in ws:
  for cell in row:
   if isinstance(cell.value,str):cell.data_type='s'
 ws.freeze_panes='F2' if 'gene' in df.columns else 'C2' if name=='State_summary' else 'A2'
 ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85 if 'gene' in df.columns else 100
 table=Table(displayName=f'MarkerTable{i+1}',ref=ws.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
 for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=44
 for j,col in enumerate(df.columns,1):
  width=12 if col in ['rank','cluster_id'] else 34 if col=='cell_state' else 23 if col=='gene_id' else 18
  if col in ['sample_consistency','specificity_flag','cohort_caution','caution']:width=28
  if col=='top10_markers':width=70
  ws.column_dimensions[get_column_letter(j)].width=width
  if pd.api.types.is_numeric_dtype(df[col]):
   fmt='0' if pd.api.types.is_integer_dtype(df[col]) else '0.0%' if col=='sample_positive_fraction' else '0.00'
   for cells in ws.iter_rows(min_row=2,min_col=j,max_col=j):cells[0].number_format=fmt
  if col in ['log2fc_mean','detection_delta_pct_points','sample_median_log2fc']:
   letter=get_column_letter(j);ws.conditional_formatting.add(f'{letter}2:{letter}{ws.max_row}',ColorScaleRule(start_type='min',start_color='E8F1FA',end_type='max',end_color='63BE7B'))
 if name=='Read_me':
  ws.column_dimensions['A'].width=33;ws.column_dimensions['B'].width=115
  for row in ws.iter_rows(min_row=2):
   row[1].alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row[0].row].height=48
 if name=='State_summary':
  for row in ws.iter_rows(min_row=2):
   row[1].hyperlink='#'+"'"+str(row[1].value)[:31]+"'!A1";row[1].font=Font(color='0563C1',underline='single')
 if 'sample_consistency' in df.columns:
  ix=df.columns.get_loc('sample_consistency')+1
  for cells in ws.iter_rows(min_row=2,min_col=ix,max_col=ix):
   if cells[0].value=='Limited samples':cells[0].fill=PatternFill('solid',fgColor='FFC7CE')
wb.save(target)
# Re-open independently, verify all stored values, dimensions, filters and names.
check=load_workbook(target,data_only=False);assert check.sheetnames==list(dfs)
for name,df in dfs.items():
 ws=check[name];assert ws.max_row==len(df)+1 and ws.max_column==len(df.columns) and len(ws.tables)==1 and ws.freeze_panes
 assert [v.value for v in ws[1]]==list(df.columns)
 for actual,expected in zip(ws.iter_rows(min_row=2),df.itertuples(index=False,name=None)):
  for cell,wanted in zip(actual,expected):
   if pd.isna(wanted):assert cell.value is None
   elif isinstance(wanted,(float,np.floating)):assert math.isclose(cell.value,float(wanted),rel_tol=1e-12,abs_tol=1e-12)
   elif wanted=='':assert cell.value in ('',None)
   else:assert cell.value==wanted
   assert cell.data_type!='f'
check.close()
result=dict(status='PASS',file=target.name,sha256=hashlib.file_digest(target.open('rb'),'sha256').hexdigest(),cells=55133,states=10,marker_rows=len(allgenes),top_rows=len(top),sheets={k:len(v) for k,v in dfs.items()},source_profile_sha256=inputs,all_exported_values_verified=True,gene_level_p_values_available=False,descriptive_only=True)
(N/'MARKER_WORKBOOK_VERIFIED.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))
